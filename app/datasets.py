"""Bounded order-file ingestion; uploaded data stays in MySQL, scoped by dataset ID."""
import csv
import hashlib
import io
import threading
import zipfile
from contextlib import nullcontext
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pymysql
from app.db import connect

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
IMPORT_SLOTS = threading.BoundedSemaphore(2)
COLUMNS = ['order_id', 'customer_name', 'order_status', 'amount', 'created_at']


def ensure_schema(connection=None):
    """Upgrade old databases once; a DML-only account can use an upgraded schema."""
    with nullcontext(connection) if connection is not None else connect() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT table_name FROM information_schema.tables "
                        "WHERE table_schema=DATABASE() AND table_name IN ('data_imports','uploaded_orders')")
            present = {next(iter(row.values())).lower() for row in cur.fetchall()}
            cur.execute("SELECT COUNT(*) AS n FROM information_schema.columns "
                        "WHERE table_schema=DATABASE() AND table_name='quality_rules' "
                        "AND column_name='dataset_id'")
            if present == {'data_imports', 'uploaded_orders'} and cur.fetchone()['n']:
                return
            # The DB lock serializes upgrades across application workers.
            cur.execute("SELECT GET_LOCK('dq_upload_schema_v1', 10) AS acquired")
            if cur.fetchone()['acquired'] != 1:
                raise ValueError('数据接入正在初始化，请稍后重试')
            try:
                cur.execute('''CREATE TABLE IF NOT EXISTS data_imports (
                    id BIGINT AUTO_INCREMENT PRIMARY KEY, name VARCHAR(200) NOT NULL,
                    source_type VARCHAR(8) NOT NULL, sha256 CHAR(64) NOT NULL,
                    row_count INT NOT NULL, created_at DATETIME(6) DEFAULT CURRENT_TIMESTAMP(6),
                    UNIQUE KEY uq_import_content (source_type, sha256)) ENGINE=InnoDB''')
                cur.execute('''CREATE TABLE IF NOT EXISTS uploaded_orders (
                    dataset_id BIGINT NOT NULL, order_id BIGINT NOT NULL,
                    customer_name VARCHAR(255) NULL, order_status VARCHAR(32) NULL,
                    amount DECIMAL(10,2) NULL, created_at DATETIME NULL,
                    PRIMARY KEY (dataset_id, order_id),
                    FOREIGN KEY (dataset_id) REFERENCES data_imports(id)) ENGINE=InnoDB''')
                cur.execute("SELECT COUNT(*) AS n FROM information_schema.columns WHERE table_schema=DATABASE() AND table_name='quality_rules' AND column_name='dataset_id'")
                if not cur.fetchone()['n']:
                    cur.execute('ALTER TABLE quality_rules ADD COLUMN dataset_id BIGINT NULL')
            finally:
                cur.execute("SELECT RELEASE_LOCK('dq_upload_schema_v1')")
        conn.commit()


def parse_file(filename, content):
    if not content or len(content) > MAX_BYTES:
        raise ValueError('请选择非空文件，最大 5 MiB')
    kind = Path(filename or '').suffix.lower().lstrip('.')
    workbook = None
    try:
        if kind == 'csv':
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError as exc:
                raise ValueError('CSV 请另存为 UTF-8 编码') from exc
            source = csv.reader(io.StringIO(text, newline=''), strict=True)
        elif kind == 'xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError('Excel 解压后过大，最多 30 MiB')
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            if len(workbook.worksheets) != 1:
                raise ValueError('Excel 必须只含一个工作表，请拆分后上传')
            source = workbook.worksheets[0].iter_rows(values_only=True)
        else:
            raise ValueError('仅支持 .csv 和 .xlsx 文件')
        header = next(source, None)
        if not header:
            raise ValueError('文件缺少表头')
        header = [str(x).strip() if x is not None else '' for x in header]
        if len(set(header)) != len(header) or not set(COLUMNS[:3]).issubset(header) or not set(header).issubset(COLUMNS):
            raise ValueError('必需表头：order_id,customer_name,order_status；可选：amount,created_at；不允许重复或未知列')
        rows, seen = [], set()
        for line, cells in enumerate(source, 2):
            if len(rows) >= MAX_ROWS:
                raise ValueError('最多允许 20000 条订单')
            if len(cells) != len(header):
                raise ValueError(f'第 {line} 行列数与表头不一致')
            values = dict(zip(header, cells))
            try:
                raw_id = values['order_id']
                if isinstance(raw_id, bool) or not str(raw_id).isdigit():
                    raise ValueError('order_id 必须是正整数')
                order_id = int(raw_id)
                if not 1 <= order_id <= 9007199254740991 or order_id in seen:
                    raise ValueError('order_id 超出范围或在本文件中重复')
                fields = []
                for key, limit in [('customer_name', 255), ('order_status', 32)]:
                    value = values[key]
                    if value == '\\N':
                        value = None
                    if value is not None and (not isinstance(value, str) or len(value) > limit):
                        raise ValueError(f'{key} 必须是最多 {limit} 字符的文本或 NULL')
                    if kind == 'xlsx' and isinstance(value, str) and value.startswith('='):
                        raise ValueError('Excel 公式不允许导入，请粘贴为值')
                    fields.append(value)
                raw_amount = values.get('amount')
                amount = None if raw_amount in (None, '', '\\N') else Decimal(str(raw_amount))
                if amount is not None and (not amount.is_finite() or abs(amount) > Decimal('99999999.99') or amount != amount.quantize(Decimal('.01'))):
                    raise ValueError('amount 必须是最多两位小数的 DECIMAL(10,2)')
                raw_date = values.get('created_at')
                date = None if raw_date in (None, '', '\\N') else (raw_date if isinstance(raw_date, datetime) else datetime.fromisoformat(str(raw_date)))
                if date and (date.tzinfo is not None or date.year < 1000 or date.microsecond):
                    raise ValueError('created_at 使用无时区、精确到秒的日期时间，年份至少 1000')
                rows.append((order_id, *fields, amount, date))
                seen.add(order_id)
            except (ValueError, InvalidOperation, TypeError) as exc:
                raise ValueError(f'第 {line} 行：{exc}') from exc
        return {'name': Path(filename).name[:200], 'source_type': kind,
                'sha256': hashlib.sha256(content).hexdigest(), 'row_count': len(rows), 'rows': rows}
    except (zipfile.BadZipFile, csv.Error, KeyError) as exc:
        raise ValueError('文件格式损坏或内容不完整') from exc
    finally:
        if workbook:
            workbook.close()


def preview(parsed):
    return {**{k: v for k, v in parsed.items() if k != 'rows'},
            'columns': COLUMNS, 'samples': [dict(zip(COLUMNS, row)) for row in parsed['rows'][:5]]}


def get_dataset(dataset_id):
    from app.quality import MissingResource
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute('SELECT * FROM data_imports WHERE id=%s', (dataset_id,))
            result = cur.fetchone()
            if not result:
                raise MissingResource(f'数据集 ID {dataset_id} 不存在')
            cur.execute('SELECT order_id,customer_name,order_status,amount,created_at FROM uploaded_orders WHERE dataset_id=%s ORDER BY order_id LIMIT 5', (dataset_id,))
            result['samples'] = cur.fetchall()
            return result


def list_datasets(before_id=None):
    ensure_schema()
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute('SELECT * FROM data_imports WHERE id < %s ORDER BY id DESC LIMIT 100', (before_id or 9223372036854775807,))
            return cur.fetchall()


def import_dataset(parsed, expected_sha256):
    if parsed['sha256'] != expected_sha256:
        raise ValueError('文件与预览内容不同，请重新预览')
    ensure_schema()
    with connect() as conn:
        try:
            with conn.cursor() as cur:
                cur.execute('INSERT INTO data_imports(name,source_type,sha256,row_count) VALUES(%s,%s,%s,%s)',
                            (parsed['name'], parsed['source_type'], parsed['sha256'], parsed['row_count']))
                dataset_id = cur.lastrowid
                for offset in range(0, len(parsed['rows']), 500):
                    cur.executemany(
                        'INSERT INTO uploaded_orders '
                        '(dataset_id,order_id,customer_name,order_status,amount,created_at) '
                        'VALUES (%s,%s,%s,%s,%s,%s)',
                        [(dataset_id, *row) for row in parsed['rows'][offset:offset + 500]],
                    )
            conn.commit()
        except pymysql.IntegrityError as exc:
            conn.rollback()
            if exc.args[0] != 1062:
                raise
            with conn.cursor() as cur:
                cur.execute('SELECT id FROM data_imports WHERE source_type=%s AND sha256=%s', (parsed['source_type'], parsed['sha256']))
                existing = cur.fetchone()
                if not existing:
                    raise
                return {**get_dataset(existing['id']), 'reused': True}
        except Exception:
            conn.rollback()
            raise
    return {**get_dataset(dataset_id), 'reused': False}
